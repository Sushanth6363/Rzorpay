"""Unrecovered handoff report — the end of the funnel (ADR-0022).

WHY THIS EXISTS
    An automated recovery engine's most important honest output is the list of people it
    could NOT recover. Everything upstream is about deciding when to act; this is about
    admitting when to stop, and handing the remainder to a human with enough context that
    they are not starting from zero.

    Without it the engine silently gives up. An opportunity exhausts its follow-ups, the
    row quietly changes state, and nobody is told. That is the same silent-loss failure the
    retry queue and reconciliation sweep exist to prevent, one level up.

WHAT COUNTS AS "EVERYTHING FAILED"
    An opportunity qualifies when it is BOTH terminal AND unpaid:
      - follow-ups exhausted, or the recovery window elapsed, or contact budget spent, or
        the top of the escalation ladder reached, AND
      - no resolution ever recorded for it.
    Anything that resolved is excluded by definition - the whole point is that this list
    contains only genuine remaining exposure.

WHAT THE MERCHANT GETS
    Not just names and amounts. Every row carries WHAT WAS ALREADY TRIED - which channels,
    how many times, when the last contact was, and why the engine stopped - because a
    collections agent who re-sends the same email the engine already sent three times is
    worse than useless. The point of the handoff is that the human starts where the
    machine left off.

INVARIANTS:
1. NEVER LISTS A CUSTOMER WHO PAID. A resolution excludes the row, unconditionally.
2. REPORTS EFFORT HONESTLY. Contact history comes from the ledger, so a row showing three
   attempts had three ledger entries. Nothing is inferred or rounded up.
3. NO INVENTED RECOMMENDATION. The suggested next step is derived from what actually
   happened, and says "review" when there is nothing defensible to suggest.
"""

from __future__ import annotations

import io
import json
import sqlite3
from dataclasses import dataclass, field
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

# Terminal reasons that mean "the engine gave up", as opposed to "the customer paid".
GAVE_UP_REASONS = (
    "MAX_FOLLOWUPS",
    "recovery window",
    "budget",
    "top of the ladder",
)


@dataclass
class UnrecoveredRow:
    """One customer the engine could not recover, with everything already tried."""

    customer_id: str
    email: str
    phone: str
    amount_paise: int
    event_type: str
    diagnosis: str
    first_seen_at: str
    last_contacted_at: str
    contacts_made: int
    channels_used: str
    escalated_to: str
    stopped_reason: str
    days_outstanding: int
    opportunity_id: str
    suggested_next_step: str = ""

    def to_dict(self) -> Dict[str, Any]:
        return {
            "Customer": self.customer_id,
            "Email": self.email,
            "Phone": self.phone,
            "Amount (Rs)": round(self.amount_paise / 100, 2),
            "Stream": self.event_type,
            "Why it failed": self.diagnosis,
            "First seen": self.first_seen_at[:19],
            "Last contacted": self.last_contacted_at[:19] if self.last_contacted_at else "never",
            "Contacts made": self.contacts_made,
            "Channels tried": self.channels_used,
            "Escalated to": self.escalated_to,
            "Why we stopped": self.stopped_reason,
            "Days outstanding": self.days_outstanding,
            "Suggested next step": self.suggested_next_step,
            "Opportunity ID": self.opportunity_id,
        }


def _days_between(earlier: str, later: Optional[str] = None) -> int:
    try:
        a = datetime.fromisoformat(str(earlier).replace("Z", "+00:00"))
        b = (
            datetime.fromisoformat(str(later).replace("Z", "+00:00"))
            if later else datetime.now(timezone.utc)
        )
        return max(0, (b - a).days)
    except Exception:
        return 0


def _suggest(row: UnrecoveredRow) -> str:
    """A next step derived from what actually happened. Never invented."""
    if row.contacts_made == 0:
        return "Never contacted - check contact details are valid before writing off"
    if "budget" in row.stopped_reason.lower():
        return "Contact budget spent - needs a human decision to approve further outreach"
    if row.event_type == "OVERDUE_B2B_INVOICE":
        return "B2B: confirm with accounts payable directly; check for a disputed line item"
    if row.diagnosis == "CARD_DECLINED":
        return "Customer never updated their payment method - offer an alternative method"
    if row.diagnosis == "INSUFFICIENT_FUNDS":
        return "Funds never arrived - consider a payment plan or a later retry window"
    if row.escalated_to in ("IVR_CALL", "AGENT_DIAL"):
        return "Full ladder exhausted including voice - review for write-off"
    return "Review manually"


def collect(conn: sqlite3.Connection, merchant_id: Optional[str] = None) -> List[UnrecoveredRow]:
    """Gather every terminal, unpaid opportunity with its full contact history."""
    conn.row_factory = sqlite3.Row

    # followup_queue holds the terminal state; live_resolutions holds who paid. The LEFT
    # JOIN plus IS NULL is what guarantees INV-1: a customer who paid can never appear.
    sql = """
        SELECT f.opportunity_id, f.merchant_id, f.customer_id, f.origin_event_id,
               f.event_json, f.diagnosis_code, f.last_action, f.attempt,
               f.first_seen_at, f.stop_reason, f.status
        FROM followup_queue f
        LEFT JOIN live_resolutions r ON r.entity_id = f.origin_event_id
        WHERE r.entity_id IS NULL
          AND f.status IN ('STOPPED', 'EXHAUSTED')
    """
    params: List[Any] = []
    if merchant_id:
        sql += " AND f.merchant_id = ?"
        params.append(merchant_id)

    try:
        candidates = conn.execute(sql, params).fetchall()
    except sqlite3.OperationalError:
        return []  # tables not created yet: nothing has run

    rows: List[UnrecoveredRow] = []
    for c in candidates:
        reason = str(c["stop_reason"] or "")
        # A row stopped because the money arrived is not a failure. Belt and braces
        # alongside the join above.
        if not any(g.lower() in reason.lower() for g in GAVE_UP_REASONS):
            continue

        try:
            event = json.loads(c["event_json"] or "{}")
        except Exception:
            event = {}

        history = conn.execute(
            """
            SELECT action_type, status, COALESCE(resolved_at, attempted_at, created_at) AS at
            FROM contact_ledger
            WHERE merchant_id = ? AND customer_id = ?
            ORDER BY at;
            """,
            (c["merchant_id"], c["customer_id"]),
        ).fetchall()

        confirmed = [h for h in history if h["status"] in ("EXECUTED", "RECONCILED_DELIVERED")]
        channels = sorted({h["action_type"] for h in confirmed})
        last_at = confirmed[-1]["at"] if confirmed else ""

        row = UnrecoveredRow(
            customer_id=str(c["customer_id"]),
            email=str(event.get("email") or event.get("customer_email") or ""),
            phone=str(event.get("phone") or event.get("contact") or ""),
            amount_paise=int(event.get("amount_paise") or 0),
            event_type=str(event.get("event_type") or "UNKNOWN"),
            diagnosis=str(c["diagnosis_code"] or "UNKNOWN"),
            first_seen_at=str(c["first_seen_at"] or ""),
            last_contacted_at=str(last_at or ""),
            contacts_made=len(confirmed),
            channels_used=", ".join(channels) if channels else "none confirmed",
            escalated_to=str(c["last_action"] or "none"),
            stopped_reason=reason,
            days_outstanding=_days_between(str(c["first_seen_at"] or "")),
            opportunity_id=str(c["opportunity_id"]),
        )
        row.suggested_next_step = _suggest(row)
        rows.append(row)

    rows.sort(key=lambda r: r.amount_paise, reverse=True)  # biggest exposure first
    return rows


def summarise(rows: List[UnrecoveredRow]) -> Dict[str, Any]:
    total = sum(r.amount_paise for r in rows)
    by_stream: Dict[str, Dict[str, Any]] = {}
    by_reason: Dict[str, int] = {}
    for r in rows:
        s = by_stream.setdefault(r.event_type, {"count": 0, "paise": 0})
        s["count"] += 1
        s["paise"] += r.amount_paise
        by_reason[r.diagnosis] = by_reason.get(r.diagnosis, 0) + 1
    return {
        "customers": len(rows),
        "total_unrecovered_paise": total,
        "total_unrecovered_rupees": round(total / 100, 2),
        "contacts_spent": sum(r.contacts_made for r in rows),
        "by_stream": by_stream,
        "by_reason": by_reason,
    }


def build_workbook(rows: List[UnrecoveredRow], merchant_id: str = "all merchants") -> bytes:
    """Produce a formatted .xlsx a collections team can work from directly."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    summary = summarise(rows)
    wb = Workbook()

    # --- Sheet 1: the summary a manager reads first -------------------------------------
    ws = wb.active
    ws.title = "Summary"
    header_font = Font(bold=True, size=13)
    ws["A1"] = "Unrecovered Revenue - Handoff Report"
    ws["A1"].font = Font(bold=True, size=15)
    ws["A2"] = f"Merchant: {merchant_id}"
    ws["A3"] = f"Generated: {datetime.now(timezone.utc).isoformat()[:19]} UTC"
    ws["A4"] = (
        "These customers were NOT recovered by the automated engine. Every one has been "
        "contacted as far as policy allows, or stopped for the stated reason."
    )
    ws["A4"].alignment = Alignment(wrap_text=True)

    ws["A6"] = "Customers unrecovered"
    ws["B6"] = summary["customers"]
    ws["A7"] = "Total still outstanding (Rs)"
    ws["B7"] = summary["total_unrecovered_rupees"]
    ws["A8"] = "Contacts already spent on them"
    ws["B8"] = summary["contacts_spent"]
    for cell in ("A6", "A7", "A8"):
        ws[cell].font = Font(bold=True)
    ws["B7"].number_format = '#,##0.00'

    ws["A10"] = "By stream"
    ws["A10"].font = header_font
    ws.append([])
    r = 11
    ws.cell(r, 1, "Stream").font = Font(bold=True)
    ws.cell(r, 2, "Customers").font = Font(bold=True)
    ws.cell(r, 3, "Amount (Rs)").font = Font(bold=True)
    for stream, data in sorted(summary["by_stream"].items(), key=lambda kv: -kv[1]["paise"]):
        r += 1
        ws.cell(r, 1, stream)
        ws.cell(r, 2, data["count"])
        ws.cell(r, 3, round(data["paise"] / 100, 2)).number_format = '#,##0.00'

    r += 2
    ws.cell(r, 1, "By failure reason").font = header_font
    r += 1
    ws.cell(r, 1, "Reason").font = Font(bold=True)
    ws.cell(r, 2, "Customers").font = Font(bold=True)
    for reason, count in sorted(summary["by_reason"].items(), key=lambda kv: -kv[1]):
        r += 1
        ws.cell(r, 1, reason)
        ws.cell(r, 2, count)

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 18

    # --- Sheet 2: the working list ------------------------------------------------------
    ws2 = wb.create_sheet("Unrecovered customers")
    if rows:
        headers = list(rows[0].to_dict().keys())
        fill = PatternFill("solid", fgColor="4338CA")
        for col, name in enumerate(headers, start=1):
            cell = ws2.cell(1, col, name)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        for row_idx, row in enumerate(rows, start=2):
            for col, value in enumerate(row.to_dict().values(), start=1):
                ws2.cell(row_idx, col, value)
        for col, name in enumerate(headers, start=1):
            width = max(len(str(name)), *(len(str(r.to_dict()[name])) for r in rows))
            ws2.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 12), 46)
        ws2.freeze_panes = "A2"
        ws2.auto_filter.ref = ws2.dimensions
    else:
        ws2["A1"] = "No unrecovered customers. Everything the engine pursued either "
        ws2["A2"] = "resolved or is still in progress."

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
